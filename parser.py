import re
import io
import docx
import pypdf

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    filename_lower = filename.lower()
    if filename_lower.endswith('.docx'):
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    full_text.append(row_text)
        return "\n".join(full_text)
    elif filename_lower.endswith('.pdf'):
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        full_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                full_text.append(t)
        return "\n".join(full_text)
    else:
        try:
            return file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            return file_bytes.decode('latin-1', errors='ignore')

def parse_question_paper(text: str):
    """
    Parses MCQ question paper text into structured JSON models while:
    1. Preserving original multi-line questions and stem formatting.
    2. Respecting sub-statements (1. ... 2. ...) inside question stems.
    3. Splitting inline options like (a) 1(b) 4/3(c) 5/3(d) 2.
    4. Deduplicating repeated option lines.
    5. Preserving KaTeX / LaTeX math equations.
    6. Filtering out paper headers/titles before Question 1.
    """
    lines = text.replace('\r\n', '\n').split('\n')
    
    # Question start pattern: "1.", "1)", "Q1", "Q.1", "1:", "Question 1"
    q_start_regex = re.compile(
        r'^\s*(?:Question\s+|Q\.?\s*|)(\d+)[\.\):\s]+(.*)$', re.IGNORECASE
    )

    raw_blocks = []
    current_block = []

    expected_q_num = 1
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if current_block:
                current_block.append("")
            continue

        match = q_start_regex.match(line)
        if match:
            q_num_val = int(match.group(1))
            is_new_q = False
            if q_num_val == expected_q_num:
                is_new_q = True
            elif q_num_val == expected_q_num + 1 and current_block:
                block_str = "\n".join(current_block)
                if re.search(r'\([a-dA-D]\)|[a-dA-D][\.\)]', block_str):
                    is_new_q = True
            
            if is_new_q:
                if current_block:
                    raw_blocks.append("\n".join(current_block).strip())
                    current_block = []
                expected_q_num = q_num_val + 1
        
        current_block.append(line)
    
    if current_block:
        raw_blocks.append("\n".join(current_block).strip())

    questions = []
    q_counter = 1
    for block in raw_blocks:
        q_obj = parse_single_question_block(block, default_num=q_counter)
        # If it's a preamble/header before Q1 with no options, ignore it
        if q_counter == 1 and len(q_obj["options"]) == 0 and ("question" not in block.lower() or len(block) < 150):
            continue
        
        # Renumber to ensure 1..N sequence
        q_obj["id"] = f"q{q_counter}"
        q_obj["number"] = q_counter
        questions.append(q_obj)
        q_counter += 1

    return questions

def parse_single_question_block(block: str, default_num: int):
    lines = block.split('\n')
    
    first_line = lines[0]
    q_num_match = re.match(r'^\s*(?:Question\s+|Q\.?\s*|)(\d+)[\.\):\s]+(.*)$', first_line, re.IGNORECASE)
    
    if q_num_match:
        q_number = int(q_num_match.group(1))
        lines[0] = q_num_match.group(2)
    else:
        q_number = default_num

    stem_lines = []
    options = []
    seen_option_keys = set()
    seen_option_texts = set()
    parse_warnings = []

    def process_line_for_inline_options(line_str):
        matches = list(re.finditer(r'(?:\(([a-dA-D])\)|([a-dA-D])[\.\)])\s*([^\(\)]*?)(?=(?:\(([a-dA-D])\)|[a-dA-D][\.\)])|$)', line_str))
        if len(matches) >= 2:
            extracted = []
            for m in matches:
                key = (m.group(1) or m.group(2)).lower()
                val = m.group(3).strip()
                extracted.append((key, val))
            return extracted
        return None

    in_options = False

    for line in lines:
        line_stripped = line.strip()
        if not line_stripped:
            if not in_options and stem_lines and stem_lines[-1] != "":
                stem_lines.append("")
            continue

        inline_result = process_line_for_inline_options(line_stripped)
        if inline_result and len(inline_result) >= 2:
            in_options = True
            for k, val in inline_result:
                if k not in seen_option_keys and val not in seen_option_texts:
                    options.append({"key": k, "text": val})
                    seen_option_keys.add(k)
                    seen_option_texts.add(val)
            continue

        opt_match = re.match(r'^\s*(?:\(([a-dA-D])\)|([a-dA-D])[\.\)]|([a-dA-D])\))\s+(.*)$', line_stripped)
        
        if not opt_match:
            num_opt_match = re.match(r'^\s*(?:\(([1-4])\)|([1-4])\))\s+(.*)$', line_stripped)
            if num_opt_match and (in_options or "which of the above" in block.lower()):
                k_num = num_opt_match.group(1) or num_opt_match.group(2)
                key_map = {'1': 'a', '2': 'b', '3': 'c', '4': 'd'}
                opt_match = (key_map[k_num], num_opt_match.group(3))

        if opt_match:
            if isinstance(opt_match, tuple):
                k, val = opt_match
            else:
                k = (opt_match.group(1) or opt_match.group(2) or opt_match.group(3)).lower()
                val = opt_match.group(4).strip()
            
            in_options = True
            if k not in seen_option_keys:
                options.append({"key": k, "text": val})
                seen_option_keys.add(k)
                seen_option_texts.add(val)
            continue

        if not in_options:
            stem_lines.append(line)
        else:
            if options:
                options[-1]["text"] += "\n" + line_stripped

    stem_text = "\n".join(stem_lines).strip()

    if len(options) < 2:
        parse_warnings.append(f"Question {q_number} has fewer than 2 options parsed ({len(options)} found).")
    elif len(options) > 6:
        parse_warnings.append(f"Question {q_number} has unusually high option count ({len(options)}).")

    # Auto-classify topic based on stem keywords
    def detect_topic(text_val):
        t = text_val.lower()
        if any(k in t for k in ['remainder', 'divided by', 'divisible']):
            return "Divisibility & Remainder Theorem"
        elif any(k in t for k in ['unit digit', 'last digit']):
            return "Unit Digit & Cyclicity"
        elif any(k in t for k in ['prime', 'prime factors', 'hcf', 'lcm']):
            return "Primes, HCF & LCM"
        elif any(k in t for k in ['series', 'fraction', 'sum of', 'sqrt', 'square root', '\\frac']):
            return "Algebraic Series & Fractions"
        elif any(k in t for k in ['factor', 'highest power', 'zeros', '!']):
            return "Factors & Factorials"
        return "Number System Fundamentals"

    return {
        "id": f"q{q_number}",
        "number": q_number,
        "rawText": block,
        "stem": stem_text,
        "topic": detect_topic(stem_text),
        "options": options,
        "correctAnswer": None,
        "explanation": None,
        "marks": 1,
        "negativeMarks": 0,
        "parseWarnings": parse_warnings
    }

def parse_answer_key(text: str):
    answers = {}
    explanations = {}
    lines = text.replace('\r\n', '\n').split('\n')

    line_pattern = re.compile(
        r'^\s*(?:Q\.?\s*|Question\s*|)(\d+)[\.\)\:\s\-]+(?:\()?([a-dA-D1-4])(?:\))?\b(?:\s*[\:\-\|]\s*|\s+)?(.*)$', re.IGNORECASE
    )
    
    inline_pattern = re.compile(r'(\d+)\s*[\:\-\(]\s*([a-dA-D1-4])(?:\))?')

    for line in lines:
        line_str = line.strip()
        if not line_str:
            continue

        match = line_pattern.match(line_str)
        if match:
            q_num = int(match.group(1))
            ans_raw = match.group(2).lower()
            key_map = {'1': 'a', '2': 'b', '3': 'c', '4': 'd'}
            ans_key = key_map.get(ans_raw, ans_raw)
            exp = match.group(3).strip() if match.group(3) else ""
            exp = re.sub(r'^(?:Explanation\s*\:\s*|[\-\|]\s*)', '', exp, flags=re.IGNORECASE).strip()
            
            answers[q_num] = ans_key
            if exp:
                explanations[q_num] = exp
            continue

        inline_matches = list(inline_pattern.finditer(line_str))
        if inline_matches:
            for m in inline_matches:
                q_num = int(m.group(1))
                ans_raw = m.group(2).lower()
                key_map = {'1': 'a', '2': 'b', '3': 'c', '4': 'd'}
                ans_key = key_map.get(ans_raw, ans_raw)
                answers[q_num] = ans_key

    return {
        "answers": answers,
        "explanations": explanations
    }
